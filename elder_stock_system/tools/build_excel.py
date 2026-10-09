"""生成 01_埃尔德选股评分模型.xlsx(全部为 Excel 公式,可粘贴自己的 OHLCV 数据)。"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from elder_lib import Params, to_weekly  # noqa: E402

OUT = Path(__file__).parent.parent / "step1" / "01_埃尔德选股评分模型.xlsx"
CSV = Path(__file__).parent.parent / "step1" / "assets" / "sample_ohlcv.csv"

P = Params()
NAVY, TEAL, GOLD, GREY = "1F3A5F", "2A7F8E", "E8B04B", "F2F4F7"
HFILL = PatternFill("solid", fgColor=NAVY)
GFILL = PatternFill("solid", fgColor=TEAL)
INFILL = PatternFill("solid", fgColor="FFF6D6")      # 输入单元格(浅黄)
OUTFILL = PatternFill("solid", fgColor="E3F1F4")     # 输出单元格(浅蓝)
thin = Side(style="thin", color="C9CED6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WHITE = Font(color="FFFFFF", bold=True, name="Microsoft YaHei", size=10)
BASE = Font(name="Microsoft YaHei", size=10)
BOLD = Font(name="Microsoft YaHei", size=10, bold=True)


def hdr(ws, row, col, text, fill=HFILL, width=None):
    c = ws.cell(row=row, column=col, value=text)
    c.font, c.fill, c.border = WHITE, fill, BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if width:
        ws.column_dimensions[CL(col)].width = width
    return c


def put(ws, ref, value, font=BASE, fill=None, fmt=None, align=None, border=True):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BORDER
    c.alignment = align or Alignment(vertical="center", wrap_text=True)
    return c


# ───────────────────────────── 数据 ─────────────────────────────
daily = pd.read_csv(CSV, parse_dates=["date"])
N = len(daily)
R0, R1 = 5, 5 + N - 1                      # 日线数据首末行
WEEKS = len(to_weekly(daily))
WK_ROWS = 80                                # 周线预留行数
W0, W1 = 5, 5 + WK_ROWS - 1

wb = Workbook()
ws_help = wb.active
ws_help.title = "说明"
ws_par = wb.create_sheet("参数")
ws_sc = wb.create_sheet("评分卡")
ws_d = wb.create_sheet("日线")
ws_w = wb.create_sheet("周线")
ws_pos = wb.create_sheet("仓位与风险")
ws_tr = wb.create_sheet("交易评级")
ws_mk = wb.create_sheet("市场温度")
ws_ch = wb.create_sheet("图表")

# ───────────────────────────── 参数 ─────────────────────────────
params = [
    ("趋势与价值区间", None, None, None),
    ("p_EmaF", "快速EMA周期(价值区间上沿)", P.ema_fast, "第22节:与慢线保持约 2:1(13/26)"),
    ("p_EmaS", "慢速EMA周期(趋势线/价值区间下沿)", P.ema_slow, "第22节:26≈半年周数,22≈月交易日"),
    ("p_MacdF", "MACD 快线周期", P.macd_fast, "第23节(12-26-9)"),
    ("p_MacdS", "MACD 慢线周期", P.macd_slow, "第23节"),
    ("p_MacdSig", "MACD 信号线周期", P.macd_sig, "第23节"),
    ("p_AdxN", "趋向系统(ADX)平滑周期", P.adx_n, "第24节(13)"),
    ("波动率、通道、止损", None, None, None),
    ("p_AtrN", "ATR 平均天数", P.atr_n, "第24节(13日)"),
    ("p_ChPct", "通道系数(占慢EMA比例)", P.channel_pct, "第22/41节:调到覆盖近100根K线约95%;A股日线建议先用6%,按个股校准"),
    ("p_PenLb", "平均下跌穿透回溯期(根)", P.pen_lookback, "第39节:回看4-6周日线"),
    ("p_StopLb", "尼克止损回溯期(根)", P.stop_lookback, "第54节:近期次低点"),
    ("p_StopAtr", "止损距入场最少 ATR 倍数", P.stop_min_atr, "第24/54节:至少1倍ATR,2倍更稳"),
    ("震荡指标", None, None, None),
    ("p_FiS", "强力指数短周期", P.fi_short, "第30节(2日EMA,择时)"),
    ("p_FiL", "强力指数长周期", P.fi_long, "第30节(13日EMA,趋势)"),
    ("p_FiLb", "强力指数\"几周内新低\"回溯期(根)", P.fi_newlow_lb, "第39节:2日FI不能创数周新低"),
    ("p_StochN", "随机指标周期", P.stoch_n, "第26节(5)"),
    ("p_StochSm", "随机指标平滑", P.stoch_smooth, "第26节(3)"),
    ("背离识别", None, None, None),
    ("p_GapMin", "背离两谷最小间隔(根)", P.divg_min_gap, "第23节:20-40根最佳;此处放宽为10-60"),
    ("p_GapMax", "背离两谷最大间隔(根)", P.divg_max_gap, "同上"),
    ("评级与风险", None, None, None),
    ("p_RRmin", "最小盈亏比", P.rr_min, "第38/53节:至少2:1"),
    ("p_GradeA", "A级评分阈值", P.grade_a, "设计假设,待第三步回测校准"),
    ("p_GradeB", "B级评分阈值", P.grade_b, "设计假设,待第三步回测校准"),
    ("p_Equity", "交易账户资金(元)", 1000000, "第50节:仅指交易账户,不含储蓄/退休金"),
    ("p_RiskTrade", "单笔风险比例(默认1%)", P.risk_trade, "第50节:专业人士常用≤1%"),
    ("p_RiskCap", "单笔风险上限(2%法则)", P.risk_trade_cap, "第50节:绝对上限2%"),
    ("p_RiskMonth", "月度总风险上限(6%法则)", P.risk_month_cap, "第51节:已亏损+持仓风险≥6%停止开新仓"),
    ("p_PosCap", "单只股票最大仓位占比", 0.20, "设计假设:A股T+1、涨跌停缺口会放大实际亏损,限制集中度"),
    ("p_Lot", "最小交易单位(股)", P.lot, "A股:100股/手"),
]
ws_par["A1"].value = "参数表(浅黄色单元格可修改,全部公式引用此表)"
ws_par["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color=NAVY)
for j, (t, w) in enumerate([("名称(定义名)", 16), ("含义", 36), ("取值", 14), ("来源 / 依据", 62)], start=1):
    hdr(ws_par, 3, j, t, width=w)
row = 4
for name, label, val, src in params:
    if label is None:
        c = ws_par.cell(row=row, column=1, value=name)
        for j in range(1, 5):
            ws_par.cell(row=row, column=j).fill = GFILL
        c.font = WHITE
    else:
        put(ws_par, f"A{row}", name, BOLD)
        put(ws_par, f"B{row}", label)
        fmt = "0.0%" if name in ("p_ChPct", "p_RiskTrade", "p_RiskCap", "p_RiskMonth", "p_PosCap") else ("#,##0" if name == "p_Equity" else "General")
        put(ws_par, f"C{row}", val, BOLD, INFILL, fmt)
        put(ws_par, f"D{row}", src)
        wb.defined_names[name] = DefinedName(name, attr_text=f"'参数'!$C${row}")
    row += 1
ws_par.freeze_panes = "A4"

# ───────────────────────────── 日线(数据 + 指标) ─────────────────────────────
D = {}
cols = []   # (key, header, group, width, fmt)


def col(key, header, group, width=10, fmt="0.00"):
    cols.append((key, header, group, width, fmt))
    D[key] = CL(len(cols))


col("date", "日期", "行情(粘贴区)", 11, "yyyy-mm-dd")
col("open", "开盘", "行情(粘贴区)")
col("high", "最高", "行情(粘贴区)")
col("low", "最低", "行情(粘贴区)")
col("close", "收盘", "行情(粘贴区)")
col("volume", "成交量", "行情(粘贴区)", 12, "#,##0")
for k, h, g, w, f in [
    ("wk", "周键(周一)", "趋势", 11, "yyyy-mm-dd"),
    ("ema_f", "EMA快", "趋势", 9, "0.00"), ("ema_s", "EMA慢", "趋势", 9, "0.00"),
    ("e12", "EMA12", "MACD", 9, "0.000"), ("e26", "EMA26", "MACD", 9, "0.000"),
    ("macd", "MACD线", "MACD", 9, "0.000"), ("sig", "信号线", "MACD", 9, "0.000"), ("hist", "MACD柱", "MACD", 9, "0.000"),
    ("f_up", "EMA快↑", "动力系统", 7, "0"), ("f_dn", "EMA快↓", "动力系统", 7, "0"),
    ("h_up", "柱↑", "动力系统", 6, "0"), ("h_dn", "柱↓", "动力系统", 6, "0"),
    ("imp", "动力系统(1绿/-1红/0蓝)", "动力系统", 10, "0"), ("impt", "颜色", "动力系统", 6, "General"),
    ("fi1", "强力指数", "强力指数", 12, "#,##0"), ("fi2", "FI-2日EMA", "强力指数", 12, "#,##0"), ("fi13", "FI-13日EMA", "强力指数", 12, "#,##0"),
    ("tr", "真实波幅TR", "波动率/ADX", 9, "0.000"), ("atr", "ATR", "波动率/ADX", 9, "0.000"),
    ("pdm", "+DM", "波动率/ADX", 8, "0.000"), ("mdm", "-DM", "波动率/ADX", 8, "0.000"),
    ("str", "平滑TR", "波动率/ADX", 9, "0.000"), ("spdm", "平滑+DM", "波动率/ADX", 9, "0.000"), ("smdm", "平滑-DM", "波动率/ADX", 9, "0.000"),
    ("pdi", "+DI", "波动率/ADX", 8, "0.0"), ("mdi", "-DI", "波动率/ADX", 8, "0.0"), ("dx", "DX", "波动率/ADX", 8, "0.0"), ("adx", "ADX", "波动率/ADX", 8, "0.0"),
    ("ch_up", "通道上轨", "通道", 9, "0.00"), ("ch_dn", "通道下轨", "通道", 9, "0.00"),
    ("au1", "ATR+1", "通道", 8, "0.00"), ("au2", "ATR+2", "通道", 8, "0.00"), ("au3", "ATR+3", "通道", 8, "0.00"),
    ("ad1", "ATR-1", "通道", 8, "0.00"), ("ad2", "ATR-2", "通道", 8, "0.00"), ("ad3", "ATR-3", "通道", 8, "0.00"),
    ("hh", "N日最高", "随机指标", 8, "0.00"), ("ll", "N日最低", "随机指标", 8, "0.00"), ("raw", "原始%K", "随机指标", 8, "0.0"),
    ("sk", "慢%K", "随机指标", 8, "0.0"), ("sd", "慢%D", "随机指标", 8, "0.0"),
    ("pen", "下穿EMA快深度", "挂单/止损", 9, "0.000"), ("avgpen", "平均下跌穿透", "挂单/止损", 9, "0.000"),
    ("enext", "明日EMA估计", "挂单/止损", 9, "0.000"), ("entry", "买入挂单价", "挂单/止损", 9, "0.00"),
    ("nic", "近期次低点", "挂单/止损", 9, "0.00"), ("stop", "止损价", "挂单/止损", 9, "0.00"),
    ("target", "目标价(上轨)", "挂单/止损", 9, "0.00"), ("rr", "盈亏比", "挂单/止损", 8, "0.00"),
    ("neg", "柱<0", "背离", 6, "0"), ("rh", "当前区段柱谷", "背离", 9, "0.000"), ("rl", "当前区段价低", "背离", 9, "0.00"),
    ("rb", "谷位置", "背离", 7, "0"), ("ph", "前一谷柱值", "背离", 9, "0.000"), ("pl", "前一谷价低", "背离", 9, "0.00"),
    ("pb", "前谷位置", "背离", 7, "0"), ("gap", "两谷间隔", "背离", 7, "0"),
    ("bdiv", "牛市背离", "背离", 7, "0"), ("fbrk", "向下假突破", "背离", 7, "0"), ("dstr", "强背离", "背离", 7, "0"),
    ("fi2nl", "FI2非新低", "信号", 7, "0"), ("fi2nm", "FI2负值均值", "信号", 12, "#,##0"), ("extr", "极端超卖反弹", "信号", 8, "0"),
    ("vma5", "5日均量", "信号", 12, "#,##0"),
    ("wk_up", "上周EMA慢↑", "周线状态(上一完整周)", 8, "0"), ("wk_hup", "上周柱↑", "周线状态(上一完整周)", 8, "0"),
    ("wk_imp", "上周动力", "周线状态(上一完整周)", 8, "0"),
    ("sa", "形态A回调到价值区", "评分", 9, "0"), ("sb", "形态B背离", "评分", 8, "0"), ("sc", "形态C极端", "评分", 8, "0"),
    ("s_tr", "趋势25", "评分", 7, "0"), ("s_pu", "回调25", "评分", 7, "0"), ("s_si", "形态25", "评分", 7, "0"),
    ("s_vo", "量能10", "评分", 7, "0"), ("s_rr", "盈亏15", "评分", 7, "0"), ("score", "总分", "评分", 7, "0"),
    ("v1", "否决:周趋势", "评分", 7, "0"), ("v2", "否决:周红", "评分", 7, "0"), ("v3", "否决:日红", "评分", 7, "0"),
    ("v4", "否决:无形态", "评分", 7, "0"), ("v5", "否决:RR<2", "评分", 7, "0"), ("grade", "评级", "评分", 7, "General"),
]:
    col(k, h, g, w, f)

n = lambda key, r: f"{D[key]}{r}"            # noqa: E731  当前行单元格
K = lambda size: f"MIN({size},ROW()-{R0 - 1})"   # noqa: E731  窗口(含当前行)已有根数
KP = lambda size: f"MIN({size},ROW()-{R0})"      # noqa: E731  不含当前行


def win(key, r, size, incl=True):
    k = K(size) if incl else KP(size)
    if incl:
        return f"OFFSET({n(key, r)},1-{k},0,{k},1)"
    return f"OFFSET({n(key, r)},-{k},0,{k},1)"


def ema_f(key_src, r, period, key_dst):
    a = f"2/({period}+1)"
    if r == R0:
        return f"={n(key_src, r)}"
    return f"={a}*{n(key_src, r)}+(1-{a})*{n(key_dst, r - 1)}"


def daily_formulas(r):
    f = {}
    first = r == R0
    c = lambda k: n(k, r)            # noqa: E731
    p = lambda k: n(k, r - 1)        # noqa: E731
    f["wk"] = f"={c('date')}-WEEKDAY({c('date')},3)"
    f["ema_f"] = ema_f("close", r, "p_EmaF", "ema_f")
    f["ema_s"] = ema_f("close", r, "p_EmaS", "ema_s")
    f["e12"] = ema_f("close", r, "p_MacdF", "e12")
    f["e26"] = ema_f("close", r, "p_MacdS", "e26")
    f["macd"] = f"={c('e12')}-{c('e26')}"
    f["sig"] = ema_f("macd", r, "p_MacdSig", "sig")
    f["hist"] = f"={c('macd')}-{c('sig')}"
    f["f_up"] = "=0" if first else f"=IF({c('ema_f')}>{p('ema_f')},1,0)"
    f["f_dn"] = "=0" if first else f"=IF({c('ema_f')}<{p('ema_f')},1,0)"
    f["h_up"] = "=0" if first else f"=IF({c('hist')}>{p('hist')},1,0)"
    f["h_dn"] = "=0" if first else f"=IF({c('hist')}<{p('hist')},1,0)"
    f["imp"] = "=0" if first else f"=IF(AND({c('f_up')}=1,{c('h_up')}=1),1,IF(AND({c('f_dn')}=1,{c('h_dn')}=1),-1,0))"
    f["impt"] = f'=IF({c("imp")}=1,"绿",IF({c("imp")}=-1,"红","蓝"))'
    f["fi1"] = "=0" if first else f"={c('volume')}*({c('close')}-{p('close')})"
    f["fi2"] = ema_f("fi1", r, "p_FiS", "fi2")
    f["fi13"] = ema_f("fi1", r, "p_FiL", "fi13")
    f["tr"] = (f"={c('high')}-{c('low')}" if first else
               f"=MAX({c('high')}-{c('low')},ABS({c('high')}-{p('close')}),ABS({c('low')}-{p('close')}))")
    f["atr"] = f"=AVERAGE({win('tr', r, 'p_AtrN')})"
    f["pdm"] = "=0" if first else f"=IF(AND({c('high')}-{p('high')}>{p('low')}-{c('low')},{c('high')}-{p('high')}>0),{c('high')}-{p('high')},0)"
    f["mdm"] = "=0" if first else f"=IF(AND({p('low')}-{c('low')}>{c('high')}-{p('high')},{p('low')}-{c('low')}>0),{p('low')}-{c('low')},0)"
    f["str"] = ema_f("tr", r, "p_AdxN", "str")
    f["spdm"] = ema_f("pdm", r, "p_AdxN", "spdm")
    f["smdm"] = ema_f("mdm", r, "p_AdxN", "smdm")
    f["pdi"] = f"=100*{c('spdm')}/{c('str')}"
    f["mdi"] = f"=100*{c('smdm')}/{c('str')}"
    f["dx"] = f"=IFERROR(100*ABS({c('pdi')}-{c('mdi')})/({c('pdi')}+{c('mdi')}),0)"
    f["adx"] = ema_f("dx", r, "p_AdxN", "adx")
    f["ch_up"] = f"={c('ema_s')}*(1+p_ChPct)"
    f["ch_dn"] = f"={c('ema_s')}*(1-p_ChPct)"
    for i in (1, 2, 3):
        f[f"au{i}"] = f"={c('ema_s')}+{i}*{c('atr')}"
        f[f"ad{i}"] = f"={c('ema_s')}-{i}*{c('atr')}"
    f["hh"] = f"=MAX({win('high', r, 'p_StochN')})"
    f["ll"] = f"=MIN({win('low', r, 'p_StochN')})"
    f["raw"] = f"=IF({c('hh')}={c('ll')},50,100*({c('close')}-{c('ll')})/({c('hh')}-{c('ll')}))"
    f["sk"] = f"=AVERAGE({win('raw', r, 'p_StochSm')})"
    f["sd"] = f"=AVERAGE({win('sk', r, 'p_StochSm')})"
    f["pen"] = f"=MAX(0,{c('ema_f')}-{c('low')})"
    cnt = f"COUNTIF({win('pen', r, 'p_PenLb')},\">0\")"
    f["avgpen"] = f"=IF({cnt}>0,SUM({win('pen', r, 'p_PenLb')})/{cnt},0.5*{c('atr')})"
    f["enext"] = f"={c('ema_f')}" if first else f"=2*{c('ema_f')}-{p('ema_f')}"
    f["entry"] = f"=ROUND(MAX({c('enext')}-{c('avgpen')},0.01),2)"
    f["nic"] = f"=IF({K('p_StopLb')}<2,{c('low')},SMALL({win('low', r, 'p_StopLb')},2))"
    raw_stop = f"MIN({c('nic')}-0.01,{c('entry')}-p_StopAtr*{c('atr')})"
    f["stop"] = f"=IF(MOD(ROUND(ROUNDDOWN({raw_stop},2)*100,0),50)=0,ROUNDDOWN({raw_stop},2)-0.01,ROUNDDOWN({raw_stop},2))"
    f["target"] = f"=ROUND({c('ch_up')},2)"
    f["rr"] = f"=IF({c('entry')}-{c('stop')}>0,({c('target')}-{c('entry')})/({c('entry')}-{c('stop')}),0)"
    # 背离
    f["neg"] = f"=IF({c('hist')}<0,1,0)"
    if first:
        f["rh"] = f'=IF({c("neg")}=1,{c("hist")},"")'
        f["rl"] = f'=IF({c("neg")}=1,{c("low")},"")'
        f["rb"] = f'=IF({c("neg")}=1,ROW(),"")'
        f["ph"] = f["pl"] = f["pb"] = '=""'
    else:
        st = f"AND({c('neg')}=1,{p('neg')}=0)"
        f["rh"] = f"=IF({st},{c('hist')},IF({c('neg')}=1,MIN({p('rh')},{c('hist')}),{p('rh')}))"
        f["rl"] = f"=IF({st},{c('low')},IF({c('neg')}=1,MIN({p('rl')},{c('low')}),{p('rl')}))"
        f["rb"] = f"=IF({st},ROW(),IF({c('neg')}=1,IF({c('hist')}<{p('rh')},ROW(),{p('rb')}),{p('rb')}))"
        f["ph"] = f"=IF({st},{p('rh')},{p('ph')})"
        f["pl"] = f"=IF({st},{p('rl')},{p('pl')})"
        f["pb"] = f"=IF({st},{p('rb')},{p('pb')})"
    f["gap"] = f'=IF(AND(ISNUMBER({c("rb")}),ISNUMBER({c("pb")})),{c("rb")}-{c("pb")},"")'
    if first:
        f["bdiv"] = f["fbrk"] = f["dstr"] = "=0"
    else:
        f["bdiv"] = (f"=IF(AND({c('neg')}=1,{c('hist')}>{p('hist')},ISNUMBER({c('ph')}),ISNUMBER({c('gap')})),"
                     f"IF(AND({c('rh')}>{c('ph')},{c('rl')}<{c('pl')},{c('gap')}>=p_GapMin,{c('gap')}<=p_GapMax),1,0),0)")
        f["fbrk"] = f"=IF(AND({c('bdiv')}=1,{c('close')}>{c('pl')}),1,0)"
        f["dstr"] = f"=IF({c('bdiv')}=1,IF({c('rh')}>={c('ph')}*0.5,1,0),0)"
    f["fi2nl"] = ("=0" if first else
                  f"=IF({c('fi2')}>MIN({win('fi2', r, 'p_FiLb-1', incl=False)}),1,0)")
    negc = f"COUNTIF({win('fi2', r, 100)},\"<0\")"
    f["fi2nm"] = f'=IF({negc}>=5,AVERAGEIF({win("fi2", r, 100)},"<0"),"")'
    f["extr"] = ("=0" if first else
                 f"=IF(ISNUMBER({c('fi2nm')}),IF(AND({c('fi2')}<5*{c('fi2nm')},{c('fi2')}>{p('fi2')},{c('close')}<{c('ch_dn')}),1,0),0)")
    f["vma5"] = f"=AVERAGE({win('volume', r, 5)})"
    # 周线状态:上一"有交易的完整周"
    m = f"MATCH({c('wk')},周线!$A${W0}:$A${W1},0)"
    for key, wcol in (("wk_up", "R"), ("wk_hup", "S"), ("wk_imp", "U")):
        f[key] = f'=IFERROR(IF({m}<=1,"",INDEX(周线!${wcol}${W0}:${wcol}${W1},{m}-1)),"")'
    # 形态
    f["sa"] = (f"=IF(AND({c('wk_up')}=1,{c('imp')}<>-1,{c('fi2')}<0,{c('fi2nl')}=1,"
               f"{c('close')}<={c('ema_f')},{c('close')}>={c('ch_dn')}),1,0)")
    f["sb"] = f"={c('bdiv')}"
    f["sc"] = f"={c('extr')}"
    f["s_tr"] = (f"=IF({c('wk_up')}=1,10,0)+IF({c('wk_imp')}=1,5,IF({c('wk_imp')}=-1,0,3))+IF({c('wk_hup')}=1,5,0)"
                 f"+IF(AND({c('pdi')}>{c('mdi')},{c('adx')}>{('0' if first else p('adx'))}),5,0)")
    f["s_pu"] = (f"=IF(AND({c('close')}<={c('ema_f')},{c('close')}>={c('ema_s')}),10,"
                 f"IF(AND({c('close')}>={c('ch_dn')},{c('close')}<{c('ema_s')}),6,0))"
                 f"+IF(AND({c('fi2')}<0,{c('fi2nl')}=1),8,0)+IF({c('sk')}<30,7,IF({c('sk')}<50,3,0))")
    f["s_si"] = (f"=IF(AND({c('sa')}=1,{c('sb')}=1),25,IF({c('sb')}=1,22,IF({c('sa')}=1,15,IF({c('sc')}=1,12,0))))")
    f["s_vo"] = f"=IF({c('fi13')}>0,6,0)+IF({c('volume')}<{c('vma5')},4,0)"
    f["s_rr"] = f"=IF({c('rr')}>=3,15,IF({c('rr')}>=p_RRmin,10,IF({c('rr')}>=1.5,4,0)))"
    f["score"] = f"={c('s_tr')}+{c('s_pu')}+{c('s_si')}+{c('s_vo')}+{c('s_rr')}"
    f["v1"] = f"=IF(AND({c('wk_up')}<>1,{c('sb')}+{c('sc')}=0),1,0)"
    f["v2"] = f"=IF({c('wk_imp')}=-1,1,0)"
    f["v3"] = f"=IF({c('imp')}=-1,1,0)"
    f["v4"] = f"=IF({c('sa')}+{c('sb')}+{c('sc')}=0,1,0)"
    f["v5"] = f"=IF({c('rr')}<p_RRmin,1,0)"
    f["grade"] = (f'=IF(SUM({c("v1")}:{c("v5")})=0,IF({c("score")}>=p_GradeA,"A",IF({c("score")}>=p_GradeB,"B","C")),'
                  f'IF(AND({c("v5")}=1,SUM({c("v1")}:{c("v4")})=0,{c("score")}>=p_GradeB),"B","C"))')
    return f


ws_d["A1"].value = "日线数据与指标(A:F 粘贴自己的 OHLCV;其余列全部为公式。向下追加数据时,把最后一行公式整行下拉复制)"
ws_d["A1"].font = Font(name="Microsoft YaHei", size=12, bold=True, color=NAVY)
ws_d["A2"].value = "※ 当前为合成示例数据(DEMO),仅用于演示公式,不是真实行情。"
ws_d["A2"].font = Font(name="Microsoft YaHei", size=10, color="B00020", bold=True)
groups = {}
for i, (k, h, g, w, f) in enumerate(cols, start=1):
    hdr(ws_d, 4, i, h, width=w, fill=HFILL if g != "行情(粘贴区)" else GFILL)
    groups.setdefault(g, [i, i])[1] = i
for g, (a, b) in groups.items():
    c = ws_d.cell(row=3, column=a, value=g)
    c.font = Font(name="Microsoft YaHei", size=9, bold=True, color=NAVY)
    for j in range(a, b + 1):
        ws_d.cell(row=3, column=j).fill = PatternFill("solid", fgColor="DDE6EE")
ws_d.row_dimensions[4].height = 42
for i, rec in enumerate(daily.itertuples(index=False)):
    r = R0 + i
    vals = [rec.date.to_pydatetime(), rec.open, rec.high, rec.low, rec.close, rec.volume]
    for j, v in enumerate(vals, start=1):
        c = ws_d.cell(row=r, column=j, value=v)
        c.number_format = cols[j - 1][4]
        c.font = BASE
        c.fill = INFILL
    for k, fx in daily_formulas(r).items():
        cidx = [x[0] for x in cols].index(k) + 1
        c = ws_d.cell(row=r, column=cidx, value=fx)
        c.number_format = cols[cidx - 1][4]
        c.font = BASE
ws_d.freeze_panes = "B5"
# 条件格式
rng = lambda key: f"{D[key]}{R0}:{D[key]}{R1}"          # noqa: E731
ws_d.conditional_formatting.add(rng("impt"), FormulaRule(formula=[f'{D["impt"]}{R0}="绿"'], fill=PatternFill("solid", bgColor="B7E1B5")))
ws_d.conditional_formatting.add(rng("impt"), FormulaRule(formula=[f'{D["impt"]}{R0}="红"'], fill=PatternFill("solid", bgColor="F4B6B6")))
ws_d.conditional_formatting.add(rng("impt"), FormulaRule(formula=[f'{D["impt"]}{R0}="蓝"'], fill=PatternFill("solid", bgColor="B9D3F2")))
ws_d.conditional_formatting.add(rng("grade"), CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", bgColor="FFD166"), font=Font(bold=True)))
ws_d.conditional_formatting.add(rng("grade"), CellIsRule(operator="equal", formula=['"B"'], fill=PatternFill("solid", bgColor="CFE8E6")))


def DR(key):      # 日线整列绝对区间
    return f"日线!${D[key]}${R0}:${D[key]}${R1}"


# ───────────────────────────── 周线 ─────────────────────────────
WC = {k: CL(i) for i, k in enumerate(
    ["wk", "open", "high", "low", "close", "volume", "ema_f", "ema_s", "e12", "e26", "macd", "sig", "hist",
     "f_up", "f_dn", "h_up", "h_dn", "e_s_up", "hist_up", "hist_dn", "imp"], start=1)}
# 约定列:A周键 B开 C高 D低 E收 F量 G快 H慢 I e12 J e26 K macd L sig M hist N f_up O f_dn P h_up Q h_dn R e_s_up S hist_up T hist_dn U imp
ws_w["A1"].value = "周线(由日线按周一为键用公式汇总;动力系统/EMA 与日线同参数)"
ws_w["A1"].font = Font(name="Microsoft YaHei", size=12, bold=True, color=NAVY)
ws_w["A2"].value = "日线评分引用\"上一完整周\"的周线状态,避免用未收盘周造成前视偏差(书中允许用未完成周线并提示打折,此处取更保守的做法)。"
ws_w["A2"].font = BASE
heads = ["周键(周一)", "开盘", "最高", "最低", "收盘", "成交量", "EMA快", "EMA慢", "EMA12", "EMA26", "MACD", "信号", "MACD柱",
         "EMA快↑", "EMA快↓", "柱↑", "柱↓", "EMA慢↑", "柱↑(周)", "柱↓(周)", "动力系统"]
for i, h in enumerate(heads, start=1):
    hdr(ws_w, 4, i, h, width=11 if i == 1 else 9)
for r in range(W0, W1 + 1):
    a = f"$A{r}"
    g = lambda body: f'=IF({a}="","",{body})'                  # noqa: E731
    first = r == W0
    prev = lambda cc: f"{cc}{r - 1}"                           # noqa: E731
    if first:
        ws_w[f"A{r}"].value = f"=MIN({DR('wk')})"
    else:
        nxt = f'_xlfn.MINIFS({DR("wk")},{DR("wk")},">"&A{r - 1})'
        ws_w[f"A{r}"].value = f'=IF(A{r - 1}="","",IF({nxt}=0,"",{nxt}))'
    ws_w[f"B{r}"].value = g(f"INDEX({DR('open')},MATCH({a},{DR('wk')},0))")
    ws_w[f"C{r}"].value = g(f"_xlfn.MAXIFS({DR('high')},{DR('wk')},{a})")
    ws_w[f"D{r}"].value = g(f"_xlfn.MINIFS({DR('low')},{DR('wk')},{a})")
    ws_w[f"E{r}"].value = g(f"INDEX({DR('close')},MATCH({a},{DR('wk')},1))")
    ws_w[f"F{r}"].value = g(f"SUMIFS({DR('volume')},{DR('wk')},{a})")

    def ew(cc, src, per):
        al = f"2/({per}+1)"
        return g(f"E{r}") if (first and src == "E") else (g(f"{src}{r}") if first else g(f"{al}*{src}{r}+(1-{al})*{prev(cc)}"))
    ws_w[f"G{r}"].value = ew("G", "E", "p_EmaF")
    ws_w[f"H{r}"].value = ew("H", "E", "p_EmaS")
    ws_w[f"I{r}"].value = ew("I", "E", "p_MacdF")
    ws_w[f"J{r}"].value = ew("J", "E", "p_MacdS")
    ws_w[f"K{r}"].value = g(f"I{r}-J{r}")
    ws_w[f"L{r}"].value = ew("L", "K", "p_MacdSig")
    ws_w[f"M{r}"].value = g(f"K{r}-L{r}")
    if first:
        for cc in "NOPQRSTU":
            ws_w[f"{cc}{r}"].value = g("0")
    else:
        ws_w[f"N{r}"].value = g(f"IF(G{r}>G{r - 1},1,0)")
        ws_w[f"O{r}"].value = g(f"IF(G{r}<G{r - 1},1,0)")
        ws_w[f"P{r}"].value = g(f"IF(M{r}>M{r - 1},1,0)")
        ws_w[f"Q{r}"].value = g(f"IF(M{r}<M{r - 1},1,0)")
        ws_w[f"R{r}"].value = g(f"IF(H{r}>H{r - 1},1,0)")
        ws_w[f"S{r}"].value = g(f"IF(M{r}>M{r - 1},1,0)")
        ws_w[f"T{r}"].value = g(f"IF(M{r}<M{r - 1},1,0)")
        ws_w[f"U{r}"].value = g(f"IF(AND(N{r}=1,P{r}=1),1,IF(AND(O{r}=1,Q{r}=1),-1,0))")
    for cc, fm in zip("ABCDEFGHIJKLMNOPQRSTU", ["yyyy-mm-dd"] + ["0.00"] * 4 + ["#,##0"] + ["0.00"] * 2 + ["0.000"] * 5 + ["0"] * 8):
        ws_w[f"{cc}{r}"].number_format = fm
        ws_w[f"{cc}{r}"].font = BASE
ws_w.freeze_panes = "B5"
ws_w.conditional_formatting.add(f"U{W0}:U{W1}", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", bgColor="B7E1B5")))
ws_w.conditional_formatting.add(f"U{W0}:U{W1}", CellIsRule(operator="equal", formula=["-1"], fill=PatternFill("solid", bgColor="F4B6B6")))
# 日线 -> 周线列映射修正(周线 R=EMA慢↑ S=柱↑ U=动力系统)已在日线公式中固定使用 R/S/U

# ───────────────────────────── 市场温度 ─────────────────────────────
ws_mk["A1"].value = "市场温度(第0层:大盘闸门) —— 决定今天允许投入多少风险额度"
ws_mk["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color=NAVY)
for j, (t, w) in enumerate([("指标", 34), ("输入/结果", 14), ("规则(书中依据)", 70)], start=1):
    hdr(ws_mk, 3, j, t, width=w)
mk = [
    ("基准指数周线EMA慢是否向上(1是/0否)", 1, "第22/39节:先看周线大趋势", True),
    ("基准指数周线动力系统(1绿/0蓝/-1红)", 0, "第40节:红色禁止买入", True),
    ("50日均线上方股票占比(%)", 45, "第35节:跌破25%后回升到25%之上=底部信号;>75%后回落=顶部信号", True),
    ("20日新高-新低指数", 60, "第34节:>0 多方领导强;< -500 后回升至其上=\"尖峰反弹\"短线买入信号", True),
    ("尖峰反弹信号(1有/0无)", 0, "第34/54节:20日NH-NL 跌破-500后重新站上", True),
]
for i, (lab, val, rule, inp) in enumerate(mk, start=4):
    put(ws_mk, f"A{i}", lab)
    put(ws_mk, f"B{i}", val, BOLD, INFILL)
    put(ws_mk, f"C{i}", rule)
put(ws_mk, "A10", "市场温度得分(0-4)", BOLD)
put(ws_mk, "B10", "=IF(B4=1,1,0)+IF(B5>=0,1,0)+IF(B6>=40,1,0)+IF(B7>0,1,0)", BOLD, OUTFILL, "0")
put(ws_mk, "C10", "趋势向上、周线非红、广度≥40%、NH-NL>0 各 1 分(阈值为设计假设,第三步回测校准)")
put(ws_mk, "A11", "风险额度系数(应用于单笔风险%)", BOLD)
put(ws_mk, "B11", "=IF(B5=-1,IF(B8=1,0.5,0),CHOOSE(B10+1,0,0,0.5,0.75,1))", BOLD, OUTFILL, "0%")
put(ws_mk, "C11", "红色周线→0(仅\"尖峰反弹\"允许50%试探仓);得分≤1→0;2→50%;3→75%;4→100%")
put(ws_mk, "A12", "状态", BOLD)
put(ws_mk, "B12", '=IF(B11=0,"空仓观望",IF(B11<1,"降风险","正常进攻"))', BOLD, OUTFILL)
put(ws_mk, "C12", "对应书中\"只在潮流方向交易\"的纪律:大盘不配合时,最好的交易是不交易")
wb.defined_names["m_Mult"] = DefinedName("m_Mult", attr_text="'市场温度'!$B$11")

# ───────────────────────────── 仓位与风险 ─────────────────────────────
ws_pos["A1"].value = "仓位与风险(2%法则 + 6%法则 + 铁三角)"
ws_pos["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color=NAVY)
ws_pos["A3"].value = "① 铁三角计算器:股数 = 风险额度 ÷ 每股风险(买入价−止损价),向下取整到手"
ws_pos["A3"].font = BOLD
for lab, ref, val, fmt in [("账户资金", "B4", "=p_Equity", "#,##0"), ("单笔风险比例(≤2%)", "B5", "=MIN(p_RiskTrade,p_RiskCap)", "0.0%"),
                           ("买入价", "B6", 33.09, "0.00"), ("止损价", "B7", 32.23, "0.00")]:
    put(ws_pos, f"A{ref[1:]}", lab)
    put(ws_pos, ref, val, BOLD, INFILL if not str(val).startswith("=") else OUTFILL, fmt)
put(ws_pos, "A8", "A 单笔最大风险额度(元)")
put(ws_pos, "B8", "=B4*B5", BOLD, OUTFILL, "#,##0")
put(ws_pos, "A9", "B 每股风险(元)")
put(ws_pos, "B9", "=B6-B7", BOLD, OUTFILL, "0.00")
put(ws_pos, "A10", "C=A/B 最大可买股数(按手取整)")
put(ws_pos, "B10", "=IF(B9>0,FLOOR(MIN(B8/B9,B4*p_PosCap/B6),p_Lot),0)", BOLD, OUTFILL, "#,##0")
put(ws_pos, "A11", "占用资金(元)")
put(ws_pos, "B11", "=B10*B6", BOLD, OUTFILL, "#,##0")
put(ws_pos, "A12", "占账户比例")
put(ws_pos, "B12", "=B11/B4", BOLD, OUTFILL, "0.0%")
put(ws_pos, "C10", "若结果为0:单手风险已超过额度,不要交易(第50节);同时受\"单只股票最大仓位占比\"约束", BASE)
ws_pos["A14"].value = "② 6%法则:本月已实现亏损 + 所有持仓的\"止损风险\" ≥ 月初资金的 6% → 本月不得再开新仓"
ws_pos["A14"].font = BOLD
put(ws_pos, "A15", "月初账户资金")
put(ws_pos, "B15", "=p_Equity", BOLD, OUTFILL, "#,##0")
put(ws_pos, "A16", "本月已实现亏损(正数填写)")
put(ws_pos, "B16", 0, BOLD, INFILL, "#,##0")
for j, t in enumerate(["持仓", "股数", "买入价", "当前止损价", "持仓风险(元)", "占资金比例"], start=1):
    hdr(ws_pos, 18, j, t, width=[34, 16, 12, 12, 16, 14][j - 1])
sample = [("示例A(止损已上移到保本)", 3000, 20.00, 20.00), ("示例B", 2000, 15.50, 14.80), ("示例C", 1500, 30.20, 29.10)]
for i in range(6):
    r = 19 + i
    s = sample[i] if i < len(sample) else ("", None, None, None)
    for j, v in enumerate(s, start=1):
        put(ws_pos, f"{CL(j)}{r}", v, BASE, INFILL)
    put(ws_pos, f"E{r}", f'=IF(B{r}="",0,MAX(0,(C{r}-D{r})*B{r}))', BASE, OUTFILL, "#,##0")
    put(ws_pos, f"F{r}", f"=E{r}/$B$15", BASE, OUTFILL, "0.00%")
put(ws_pos, "A26", "持仓风险合计(元)", BOLD)
put(ws_pos, "E26", "=SUM(E19:E24)", BOLD, OUTFILL, "#,##0")
put(ws_pos, "A27", "月度风险占用 = 已亏损 + 持仓风险", BOLD)
put(ws_pos, "E27", "=B16+E26", BOLD, OUTFILL, "#,##0")
put(ws_pos, "F27", "=E27/B15", BOLD, OUTFILL, "0.00%")
put(ws_pos, "A28", "可用风险额度(元)= 6% × 月初资金 − 占用", BOLD)
put(ws_pos, "E28", "=MAX(0,p_RiskMonth*B15-E27)", BOLD, OUTFILL, "#,##0")
put(ws_pos, "A29", "能否开新仓", BOLD)
put(ws_pos, "E29", '=IF(E28>0,"可以(单笔仍≤2%且≤可用额度)","停止:本月风险额度用尽")', BOLD, OUTFILL)
wb.defined_names["pos_AvailRisk"] = DefinedName("pos_AvailRisk", attr_text="'仓位与风险'!$E$28")
ws_pos["A31"].value = "说明:止损上移到成本价后该笔风险为0,释放额度(书中示例:A上移保本后才能买入D)。"
ws_pos["A31"].font = BASE
ws_pos.conditional_formatting.add("E29", FormulaRule(formula=['LEFT(E29,2)="停止"'], fill=PatternFill("solid", bgColor="F4B6B6")))

# ───────────────────────────── 交易评级 ─────────────────────────────
ws_tr["A1"].value = "交易评级(第55节):买入评级、卖出评级、交易评级(≥通道高度30% = A级交易)"
ws_tr["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color=NAVY)
heads = ["交易", "买入价", "买入日最高", "买入日最低", "卖出价", "卖出日最高", "卖出日最低", "买入日通道上轨", "买入日通道下轨",
         "买入评级", "卖出评级", "交易评级", "等级", "盈亏(元/股)"]
for i, h in enumerate(heads, start=1):
    hdr(ws_tr, 3, i, h, width=14)
ws_tr.row_dimensions[3].height = 32
ex = [("书中例: ADSK", 51.77, 52.49, 51.75, 53.78, 54.49, 53.39, 53.87, 47.61)] + [("", None, None, None, None, None, None, None, None)] * 7
for i, rowv in enumerate(ex):
    r = 4 + i
    for j, v in enumerate(rowv, start=1):
        put(ws_tr, f"{CL(j)}{r}", v, BASE, INFILL, "0.00" if j > 1 else None)
    put(ws_tr, f"J{r}", f'=IF(B{r}="","",IF(C{r}>D{r},(C{r}-B{r})/(C{r}-D{r}),""))', BOLD, OUTFILL, "0%")
    put(ws_tr, f"K{r}", f'=IF(E{r}="","",IF(F{r}>G{r},(E{r}-G{r})/(F{r}-G{r}),""))', BOLD, OUTFILL, "0%")
    put(ws_tr, f"L{r}", f'=IF(OR(B{r}="",E{r}=""),"",IF(H{r}>I{r},(E{r}-B{r})/(H{r}-I{r}),""))', BOLD, OUTFILL, "0%")
    put(ws_tr, f"M{r}", f'=IF(L{r}="","",IF(L{r}>=0.3,"A",IF(L{r}>=0.15,"B","C")))', BOLD, OUTFILL)
    put(ws_tr, f"N{r}", f'=IF(OR(B{r}="",E{r}=""),"",E{r}-B{r})', BOLD, OUTFILL, "0.00")
put(ws_tr, "A13", "解读:买入评级>50%=买在当日K线下半部;卖出评级>50%=卖在上半部;交易评级=实际盈亏/通道高度,比盈亏金额更能衡量交易质量。书中ADSK例:买入97%、卖出35%、交易评级32%(A级)。")
ws_tr.merge_cells("A13:N14")
ws_tr["A13"].alignment = Alignment(wrap_text=True, vertical="top")
ws_tr.conditional_formatting.add("M4:M11", CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", bgColor="FFD166")))

# ───────────────────────────── 评分卡 ─────────────────────────────
ws_sc["A1"].value = "评分卡 —— 单只股票打分与次日挂单计划"
ws_sc["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color=NAVY)
for j, w in enumerate([22, 30, 16, 52, 10, 10], start=1):
    ws_sc.column_dimensions[CL(j)].width = w
put(ws_sc, "A3", "标的", BOLD)
put(ws_sc, "B3", "DEMO 示例股份(合成数据)", BOLD, INFILL)
put(ws_sc, "A4", "评分日期", BOLD)
put(ws_sc, "B4", f"=MAX({DR('date')})", BOLD, INFILL, "yyyy-mm-dd")
put(ws_sc, "C4", "可改为历史任意交易日回看", BASE, border=False)
put(ws_sc, "A5", "行号(内部)", BASE)
put(ws_sc, "B5", f"=MATCH(B4,{DR('date')},0)", BASE, OUTFILL, "0")
ix = lambda key: f"INDEX({DR(key)},$B$5)"          # noqa: E731
# 评分明细表
for j, t in enumerate(["维度", "项目", "当前值", "规则", "得分", "满分"], start=1):
    hdr(ws_sc, 7, j, t)
detail = [
    ("① 趋势(25)", "上周 EMA慢是否向上", f"={ix('wk_up')}", "向上=10分(第一重滤网)", f"=IF(C8=1,10,0)", 10, "0"),
    ("", "上周动力系统", f"={ix('wk_imp')}", "绿=5 / 蓝=3 / 红=0(红色禁止买入)", f"=IF(C9=1,5,IF(C9=-1,0,3))", 5, "0"),
    ("", "上周 MACD 柱斜率向上", f"={ix('wk_hup')}", "向上=5", f"=IF(C10=1,5,0)", 5, "0"),
    ("", "ADX 上升且 +DI>-DI", f"=IF(AND({ix('pdi')}>{ix('mdi')},{ix('adx')}>INDEX({DR('adx')},MAX(1,$B$5-1))),1,0)", "趋势有劲=5(第24节)", "=IF(C11=1,5,0)", 5, "0"),
    ("② 回调质量(25)", "价格相对价值区间", f'=IF(AND({ix("close")}<={ix("ema_f")},{ix("close")}>={ix("ema_s")}),"价值区间内",IF(AND({ix("close")}>={ix("ch_dn")},{ix("close")}<{ix("ema_s")}),"慢EMA下方","区间外/追高"))', "区间内=10 / 慢线下方但在通道内=6 / 其他=0(不追高)", '=IF(C12="价值区间内",10,IF(C12="慢EMA下方",6,0))', 10, "General"),
    ("", "2日强力指数<0 且非数周新低", f"=IF(AND({ix('fi2')}<0,{ix('fi2nl')}=1),1,0)", "满足=8(第二重滤网)", "=IF(C13=1,8,0)", 8, "0"),
    ("", "慢速随机指标 %K", f"={ix('sk')}", "<30=7 / <50=3", "=IF(C14<30,7,IF(C14<50,3,0))", 7, "0.0"),
    ("③ 买入形态(25)", "形态A 回调到价值区(顺势)", f"={ix('sa')}", "仅A=15", "", None, "0"),
    ("", "形态B MACD柱牛市背离/假突破", f"={ix('sb')}", "仅B=22,A+B=25", "", None, "0"),
    ("", "形态C 极端超卖反弹", f"={ix('sc')}", "仅C=12", "=IF(AND(C15=1,C16=1),25,IF(C16=1,22,IF(C15=1,15,IF(C17=1,12,0))))", 25, "0"),
    ("④ 量能(10)", "13日强力指数>0", f"=IF({ix('fi13')}>0,1,0)", "多头占优=6", "=IF(C18=1,6,0)", 6, "0"),
    ("", "今日量<5日均量(缩量回调)", f"=IF({ix('volume')}<{ix('vma5')},1,0)", "缩量=4", "=IF(C19=1,4,0)", 4, "0"),
    ("⑤ 盈亏比(15)", "盈亏比 (目标−买入)/(买入−止损)", f"={ix('rr')}", "≥3=15 / ≥参数=10 / ≥1.5=4", "=IF(C20>=3,15,IF(C20>=p_RRmin,10,IF(C20>=1.5,4,0)))", 15, "0.00"),
]
for i, (dim, item, val, rule, sc, mx, fmt) in enumerate(detail, start=8):
    put(ws_sc, f"A{i}", dim, BOLD)
    put(ws_sc, f"B{i}", item)
    put(ws_sc, f"C{i}", val, BASE, OUTFILL, fmt)
    put(ws_sc, f"D{i}", rule)
    put(ws_sc, f"E{i}", sc if sc else None, BOLD, OUTFILL if sc else None, "0")
    put(ws_sc, f"F{i}", mx)
put(ws_sc, "A21", "总分", BOLD)
put(ws_sc, "E21", "=SUM(E8:E20)", BOLD, OUTFILL, "0")
put(ws_sc, "F21", 100, BOLD)
hdr(ws_sc, 23, 1, "硬性否决(任一触发则不得买入)")
hdr(ws_sc, 23, 2, "状态")
hdr(ws_sc, 23, 3, "触发(1)")
hdr(ws_sc, 23, 4, "说明")
veto = [
    ("周趋势闸门", "v1", "形态A要求周EMA慢向上;背离/极端形态B、C不受此限(第39节)"),
    ("周线动力系统红色", "v2", "红色禁止买入(第40节)"),
    ("日线动力系统红色", "v3", "红色禁止买入(第40节)"),
    ("无有效买入形态", "v4", "A/B/C 都不满足 = 不是\"A级交易\"(第55节)"),
    ("盈亏比不足", "v5", "盈亏比必须≥参数(默认2:1)(第38/53节)"),
]
for i, (lab, key, note) in enumerate(veto, start=24):
    put(ws_sc, f"A{i}", lab)
    put(ws_sc, f"C{i}", f"={ix(key)}", BASE, OUTFILL, "0")
    put(ws_sc, f"B{i}", f'=IF(C{i}=1,"✗ 否决","✓ 通过")', BOLD, OUTFILL)
    put(ws_sc, f"D{i}", note)
put(ws_sc, "A30", "评级(A/B/C)", BOLD)
put(ws_sc, "B30", f"={ix('grade')}", Font(name="Microsoft YaHei", size=16, bold=True, color="B00020"), OUTFILL)
put(ws_sc, "D30", '=IF(B30="A","A级交易:按计划挂单",IF(B30="B","B级:仅在额度充足时以1/3风险试探","C级:放弃,翻到下一只(第55节)"))', BOLD)
hdr(ws_sc, 32, 1, "次日挂单计划")
hdr(ws_sc, 32, 2, "数值")
hdr(ws_sc, 32, 3, "")
hdr(ws_sc, 32, 4, "算法")
plan = [
    ("买入限价", f"={ix('entry')}", "明日EMA估计 − 平均下跌穿透(第39节,以折价回调买入,不追高)", "0.00"),
    ("止损价", f"={ix('stop')}", "min(近期次低点−0.01, 买入价−N×ATR),避开整数位(第54节尼克止损)", "0.00"),
    ("目标价", f"={ix('target')}", "上通道线(第53节:在价值区间下方买、通道上轨获利)", "0.00"),
    ("每股风险", "=B33-B34", "买入价−止损价", "0.00"),
    ("盈亏比", f"={ix('rr')}", "(目标−买入)/(买入−止损)", "0.00"),
    ("大盘风险系数", "=m_Mult", "来自\"市场温度\"", "0%"),
    ("风险额度(元)", "=MIN(p_Equity*MIN(p_RiskTrade,p_RiskCap)*B38,pos_AvailRisk)", "资金×单笔风险%×大盘系数,且不超过6%法则剩余额度", "#,##0"),
    ("建议股数", "=IF(AND(B30<>\"C\",B36>0),FLOOR(MIN(B39/B36,p_Equity*p_PosCap/B33),p_Lot),0)", "铁三角:风险额度÷每股风险,并受单只仓位上限约束,按手取整", "#,##0"),
    ("占用资金(元)", "=B40*B33", "", "#,##0"),
    ("占账户比例", "=B41/p_Equity", "", "0.0%"),
]
for i, (lab, fx, note, fmt) in enumerate(plan, start=33):
    put(ws_sc, f"A{i}", lab, BOLD)
    put(ws_sc, f"B{i}", fx, BOLD, OUTFILL, fmt)
    put(ws_sc, f"D{i}", note)
ws_sc["A44"].value = "提示:分数阈值、权重与市场温度系数是本方案的设计假设(非书中原文),将在第三步用 2026 年 1 月起的历史数据回测校准。"
ws_sc["A44"].font = Font(name="Microsoft YaHei", size=9, color="B00020")
ws_sc.conditional_formatting.add("B30", CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", bgColor="FFD166")))
ws_sc.conditional_formatting.add("B24:B28", FormulaRule(formula=['LEFT(B24,1)="✗"'], fill=PatternFill("solid", bgColor="F4B6B6")))

# ───────────────────────────── 图表 ─────────────────────────────
lo = max(R0, R1 - 119)
ws_ch["A1"].value = "图表(最近 120 根日线)"
ws_ch["A1"].font = Font(name="Microsoft YaHei", size=14, bold=True, color=NAVY)
lc = LineChart()
lc.title = "收盘价、EMA13/26 价值区间与通道"
lc.height, lc.width = 9.5, 24
idx = {k: [x[0] for x in cols].index(k) + 1 for k in D}
for k in ("close", "ema_f", "ema_s", "ch_up", "ch_dn"):
    lc.add_data(Reference(ws_d, min_col=idx[k], min_row=lo, max_row=R1), titles_from_data=False)
cats = Reference(ws_d, min_col=1, min_row=lo, max_row=R1)
lc.set_categories(cats)
names = ["收盘", "EMA快(13)", "EMA慢(26)", "通道上轨", "通道下轨"]
from openpyxl.chart.series import SeriesLabel  # noqa: E402
for s, nm in zip(lc.series, names):
    s.tx = SeriesLabel(v=nm)
    s.smooth = False
    s.graphicalProperties.line.width = 14000
colors = ["1F3A5F", "E8B04B", "2A7F8E", "9AA5B1", "9AA5B1"]
for s, cc in zip(lc.series, colors):
    s.graphicalProperties.line.solidFill = cc
lc.series[3].graphicalProperties.line.dashStyle = "dash"
lc.series[4].graphicalProperties.line.dashStyle = "dash"
lc.x_axis.number_format = "mm-dd"
lc.y_axis.scaling.min = float(int(daily["low"].iloc[-120:].min() * 0.9))
lc.x_axis.tickLblPos = "low"
lc.x_axis.delete = False
lc.y_axis.delete = False
ws_ch.add_chart(lc, "A3")
bc = BarChart()
bc.title = "MACD 柱状线(多空力量斜率)"
bc.height, bc.width = 7, 24
bc.add_data(Reference(ws_d, min_col=idx["hist"], min_row=lo, max_row=R1), titles_from_data=False)
bc.set_categories(cats)
bc.series[0].tx = SeriesLabel(v="MACD柱")
bc.series[0].graphicalProperties.solidFill = "2A7F8E"
bc.x_axis.number_format = "mm-dd"
bc.x_axis.tickLblPos = "low"
bc.x_axis.delete = False
bc.y_axis.delete = False
bc.legend = None
ws_ch.add_chart(bc, "A24")
sb = BarChart()
sb.type = "bar"
sb.title = "评分构成(当前评分卡)"
sb.height, sb.width = 7, 14
sb.add_data(Reference(ws_sc, min_col=5, min_row=8, max_row=20), titles_from_data=False)
sb.set_categories(Reference(ws_sc, min_col=2, min_row=8, max_row=20))
sb.series[0].graphicalProperties.solidFill = "E8B04B"
sb.legend = None
sb.x_axis.delete = False
sb.y_axis.delete = False
ws_ch.add_chart(sb, "A39")

# ───────────────────────────── 说明 ─────────────────────────────
ws_help.column_dimensions["A"].width = 24
ws_help.column_dimensions["B"].width = 110
ws_help["A1"].value = "埃尔德《以交易为生》选股评分模型(第一步:Excel 公式版)"
ws_help["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color=NAVY)
lines = [
    ("用途", "把书中的三重滤网、动力系统、价值区间/通道、MACD柱背离+假突破、强力指数、ATR、2%/6%风险法则全部落成公式,对一只股票的日线数据打分、给出次日挂单价/止损/目标/股数。"),
    ("色彩约定", "浅黄色=可输入/可改的单元格;浅蓝色=公式输出;深蓝表头=公式列。"),
    ("工作表", "参数 → 日线(粘贴OHLCV,指标全自动)→ 周线(自动汇总)→ 评分卡(看结论)→ 仓位与风险(2%/6%)→ 交易评级 → 市场温度 → 图表。"),
    ("使用步骤", "① 在\"日线\"A:F 列粘贴自己的日期/开/高/低/收/量(示例为合成数据)。② 必要时在\"参数\"调通道系数(使近100根K线约95%落在通道内)。③ 在\"市场温度\"填大盘4个读数。④ 打开\"评分卡\"看总分、否决项、挂单计划。⑤ 买入后在\"仓位与风险\"登记,监控6%额度。⑥ 平仓后在\"交易评级\"打分复盘。"),
    ("扩展数据", "行数不够时:选中日线最后一行的公式列(G列起)向下填充即可;周线预留80周。批量处理几千只股票请使用第三步的 Python 引擎(与本表同一套定义,已逐行对账)。"),
    ("信号定义", "动力系统:EMA快↑且MACD柱↑=绿(禁止卖空);EMA快↓且柱↓=红(禁止买入);否则蓝。形态A:周EMA慢向上、日线非红、2日强力指数<0且非数周新低、收盘≤EMA快且≥通道下轨。形态B:MACD柱两个负区段谷底,价格更低而柱谷更浅、间隔在参数范围内、柱线回升。形态C:价格低于通道下轨且2日强力指数深度超过平均负值的5倍后回升。"),
    ("买入价", "三重滤网第三重:明日EMA估计 − 平均下跌穿透,即在回调中以折价挂单,不追高。"),
    ("止损", "尼克止损:近期次低点下方1分,且至少离买入价 N×ATR,并避开整数/半整数价位,以免落入拥挤的止损单区域。"),
    ("目标", "上通道线(书:在价值区间买入、通道上轨附近兑现)。分批:1/3在+1ATR,1/3在+2ATR,其余在+3ATR或动力系统转蓝/红时离场。"),
    ("风险", "单笔风险=资金×风险%(默认1%,上限2%);月度已亏损+持仓风险≥6%停止开新仓;止损上移到保本即释放额度。"),
    ("免责声明", "示例数据为合成数据。权重、阈值、大盘系数为本方案设计假设,尚未经真实历史回测验证;模型仅作研究与教育用途,不构成投资建议。"),
]
for i, (k, v) in enumerate(lines, start=3):
    put(ws_help, f"A{i}", k, BOLD, PatternFill("solid", fgColor=GREY))
    put(ws_help, f"B{i}", v)
    ws_help.row_dimensions[i].height = 48 if len(v) > 100 else 30

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved", OUT, "rows", N, "cols", len(cols))
